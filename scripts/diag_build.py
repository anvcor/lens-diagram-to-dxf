#!/usr/bin/env python3
"""lens.json（diag_extract.py 的输出）→ 镜头面数据 Excel + AutoCAD DXF（+ 可选 spec.json）。

    python3 diag_build.py lens.json -o 55EVO [--glass BSC7] [--image diagram.png] [--title "55mm F1.8"]
    → 55EVO.xlsx  55EVO.dxf  [55EVO.spec.json]

Excel 版式照用户历史文件 55EVO.xlsx：面序号 / 面型 / 曲率半径 / THICKNESS / 玻璃(占位) / 标记 / 半口径 / 拟合RMS。
  STO 行曲率半径写 Infinity；末行 IMA 的 THICKNESS 用 =SUM() 公式算总长。
DXF 照用户历史文件 55EVO.dxf：mm 单位；光轴 LINE；每个面一段 ARC（圆心在光轴上）；
  平面用 LINE；每片镜片上下各一条边缘 LINE 把前后两弧端点连起来；光阑一条竖 LINE；
  外加总长/最大口径两条 DIMENSION；--image 给了就把原图按比例贴进去（像用户那样图在下、线在上）。
坐标系：首面顶点为原点，光轴为 X 轴，光线自左向右；图像 y 向下 → DXF y 向上已翻转。
spec.json：与 patent-lens-to-excel skill 的 make_zmx.py / make_seq.py 共用的格式，玻璃统一占位。
"""
import argparse, json, math, os

GLASS_NDVD = {'BSC7': (1.51633, 64.14), 'N-BK7': (1.5168, 64.17), 'S-BSL7': (1.51633, 64.14)}


def arc_points(R, semi, n=48):
    """给定带符号半径 R（圆心在面右侧为正）与半口径，返回相对顶点的 (dx, y) 点列。"""
    if R is None:
        return [(0.0, -semi), (0.0, semi)]
    r = abs(R); semi = min(semi, r * 0.999)
    pts = []
    for i in range(n + 1):
        y = -semi + 2 * semi * i / n
        sag = r - math.sqrt(r * r - y * y)
        pts.append((sag if R > 0 else -sag, y))
    return pts


def z_of(s, y):
    """面在半径 y 处的轴向位置（mm）。R>0 向左凸，顶点最靠左。"""
    R = s['R_mm']
    if R is None: return s['_zv']
    r = abs(R)
    if y >= r: y = r * 0.999999
    sag = r - math.sqrt(r * r - y * y)
    return s['_zv'] + (sag if R > 0 else -sag)


def clearance_clamp(surfaces, pxmm, x0):
    """相邻两面延伸到各自边缘时会互相穿透（边缘厚度变负），作图必须在穿透半径处截断，
    否则 CAD 里就是「口径打架」：两片镜子的轮廓交叉成 X。

    对每一对相邻面求最小的 y 使轴向间距 ≤0，把两面的作图半径都压到它之下。
    返回值写进 s['draw_semi']；s['semi_mm'] 仍是从图上量到的机械半径，不动。
    """
    for s in surfaces:
        s['_zv'] = (s['vertex'] - x0) / pxmm
        # 作图半径先取「拟合时真正吃到的半径」——图上这个面画到哪儿为止。
        # 用机械半径（元件外框半高）会让弧冲出面的实际范围，压到邻片里去。
        s['draw_semi'] = s.get('semi_fit_mm') or s['semi_mm']
    notes = []
    for i in range(len(surfaces) - 1):
        a, b = surfaces[i], surfaces[i + 1]
        if a.get('stop') or b.get('stop'): continue
        if a['semi_mm'] is None or b['semi_mm'] is None: continue
        top = min(a['draw_semi'], b['draw_semi'])
        if z_of(b, top) - z_of(a, top) > 0: continue          # 到边缘都不打架
        lo, hi = 0.0, top
        for _ in range(60):
            mid = (lo + hi) / 2
            if z_of(b, mid) - z_of(a, mid) > 0: lo = mid
            else: hi = mid
        y0 = lo * 0.995
        for s in (a, b):
            if s['draw_semi'] > y0:
                s['draw_semi'] = y0
        notes.append((i, y0, top))
    # 不要再把「同一片的前后两面」强行取齐：胶合面是两片共用的，取齐会顺着胶合链
    # 把整组都压到最小值（实测绿片从 12.09 被拖到 9.88）。半径不等时 build_dxf 画台阶即可。
    return notes


def build_xlsx(d, path, glass, title):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = '面数据'
    hdr = ['面序号', '面型', '曲率半径', 'THICKNESS', '玻璃', '标记', '半口径', '作图半径', '拟合RMS(px)', '备注']
    ws.append([title or d.get('image', '')]); ws['A1'].font = Font(bold=True, size=12)
    ws.append([f"比例 {d['pxmm']:.4f} px/mm；总长(首面顶点→末面顶点) {d['total_length_mm']} mm；玻璃为占位 {glass}，待定"])
    ws.append(hdr)
    for c in ws[3]: c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='DDEBF7')
    surfaces = d['surfaces']; first = ws.max_row + 1; k = 0
    for s in surfaces:
        if s.get('stop'):
            name = 'STO'
        else:
            k += 1; name = k
        typ = '非球面' if s.get('asph') else '球面'
        R = 'Infinity' if s['R_mm'] is None else s['R_mm']
        note = []
        if s.get('cemented'): note.append('胶合面')
        if s.get('thin_gap'): note.append('★两面在轴上几乎相接（图上两条描边在轴附近并成一条），间隔已在图纸分辨力以下；真实设计应有 0.05~0.2 mm 气隙，进软件前请自行给一个最小值')
        if s.get('asph_by_tag'): note.append('图例标为非球面，表中 R 为最佳拟合球面的近似值')
        elif s.get('asph'): note.append('拟合残差大，可能非球面/需人工核对')
        ds = s.get('draw_semi')
        if ds is not None and s['semi_mm'] is not None and ds < s['semi_mm'] - 0.05:
            note.append(f'图上这个面只画到 r={ds:.2f} mm（再往外是机械磨边或与邻片相接），作图半径按此截断')
        ws.append([name, typ, R, s['D'], glass if s['glass'] else None, s.get('tag'),
                   s['semi_mm'], None if ds is None else round(ds, 2),
                   round(s['rms'], 2) if s['rms'] else None, '；'.join(note) or None])
    last = ws.max_row
    ws.append(['IMA', '球面', None, f'=SUM(D{first}:D{last})'])
    for col, w in zip('ABCDEFGHIJ', (8, 8, 12, 12, 10, 12, 9, 9, 11, 46)): ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=first, max_row=last + 1):
        for c in row: c.alignment = Alignment(horizontal='center')
    # 元件表
    we = wb.create_sheet('元件')
    we.append(['元件', '前顶点 x(mm)', '后顶点 x(mm)', '中心厚 (mm)', '半口径 (mm)', '颜色类', '标记'])
    tags = {int(k): v for k, v in d.get('tags', {}).items()}
    for i, e in enumerate(d['elements'], 1):
        fx, bx = e['front_vertex'] / d['pxmm'], e['back_vertex'] / d['pxmm']
        x0 = d['surfaces'][0]['vertex'] / d['pxmm']
        we.append([i, round(fx - x0, 3), round(bx - x0, 3), round(bx - fx, 3), e['semi_mm'], e['color'], tags.get(e['color'])])
    wb.save(path)


def build_dxf(d, path, image=None, jpeg=True, fade=50):
    import ezdxf
    doc = ezdxf.new('R2018', setup=True); doc.units = ezdxf.units.MM; doc.header['$INSUNITS'] = 4
    msp = doc.modelspace()
    for name, color in (('AXIS', 1), ('SURF', 7), ('EDGE', 7), ('STOP', 3), ('DIM', 4), ('IMG', 8)):
        doc.layers.add(name, color=color)
    pxmm = d['pxmm']; ay = d['axis_px']; x0 = d['surfaces'][0]['vertex']
    surfaces = d['surfaces']
    total = d['total_length_mm']; semi_max = max(s['semi_mm'] or 0 for s in surfaces)
    if image:
        # 原图按比例贴进去：像素 (x, y) → ((x - x0)/pxmm, (ay - y)/pxmm)
        W, H = d['W'], d['H']
        # 外参文件名必须纯 ASCII：DXF 里写的是 UTF-8，AutoCAD 按 ANSI 码页读，
        # 中文名会变成 ?? 而解析不到图（图上只剩一行灰色占位文字）。名字带非 ASCII 就改名并拷一份。
        # AutoCAD 的光栅引擎读 PNG 常常报「不可读」（文件找得到、解不开）。用户那份能正常显示的
        # 55EVO.dxf 贴的就是 .jpg，所以这里一律转成 JPEG，并照它的写法加 `.\` 前缀。
        base = os.path.basename(image)
        stem = ''.join(c if ord(c) < 128 else '_' for c in os.path.splitext(base)[0])
        outdir = os.path.dirname(os.path.abspath(path))
        if jpeg:
            from PIL import Image as _Im
            safe = stem + '.jpg'
            im = _Im.open(image)
            if im.mode != 'RGB': im = im.convert('RGB')
            im.save(os.path.join(outdir, safe), 'JPEG', quality=92, subsampling=0)
            print(f'  贴图已转存为 {safe}（AutoCAD 读 PNG 常报「不可读」，JPEG 最稳）')
        else:
            safe = stem + os.path.splitext(base)[1]
            if safe != base:
                import shutil
                shutil.copyfile(image, os.path.join(outdir, safe))
                print(f'  贴图名含非 ASCII 字符，已另存为 {safe}')
        # ACAD_IMAGE_DICT 的键必须是「去掉路径和扩展名的纯名字」（AutoCAD 自己就是这么写的）。
        # ezdxf 默认拿整个文件名当键，于是键里带上了 `.\` —— 反斜杠在 AutoCAD 符号名里非法，
        # 注册不进字典，AutoCAD 打开时就把引用它的 IMAGE 实体整个丢掉（外部参照面板显示「不可读」）。
        idef = doc.add_image_def(filename='.\\' + safe, size_in_pixel=(W, H),
                                 name=os.path.splitext(safe)[0])
        idef.dxf.pixel_size = (1.0 / W, 1.0 / W)   # AutoCAD 的默认像素尺寸＝图宽归一化成 1 个单位
        img = msp.add_image(image_def=idef, insert=(-x0 / pxmm, (ay - H) / pxmm),
                            size_in_units=(W / pxmm, H / pxmm), dxfattribs={'layer': 'IMG'})
        # add_image 算出来的 U/V 会带 1e-19 量级的脏分量，直接写死成正交
        img.dxf.u_pixel = (1.0 / pxmm, 0.0, 0.0)
        img.dxf.v_pixel = (0.0, 1.0 / pxmm, 0.0)
        img.dxf.flags = 7          # 显示 + 未对齐时也显示 + 使用裁剪边界（照 55EVO 那份）
        img.dxf.fade = fade        # 淡显，压在下面当描图底板
        doc.set_raster_variables(frame=1, quality=1, units='mm')  # ezdxf 默认 'm'，mm 图要显式给
    msp.add_line((-5, 0), (total + 5, 0), dxfattribs={'layer': 'AXIS', 'linetype': 'CENTER'})
    # 各面
    ends = []  # 每面 (x_top, y_top, x_bot, y_bot) 供边缘线用
    for s in surfaces:
        xv = (s['vertex'] - x0) / pxmm
        if s.get('stop'):
            h = s.get('semi_mm') or (d.get('stop_semi_px') or [semi_max * pxmm * 0.9])[0] / pxmm
            msp.add_line((xv, -h), (xv, h), dxfattribs={'layer': 'STOP'})
            ends.append(None); continue
        semi = s.get('draw_semi', s['semi_mm'])
        if s['R_mm'] is None:
            msp.add_line((xv, -semi), (xv, semi), dxfattribs={'layer': 'SURF'})
            ends.append((xv, semi, xv, -semi)); continue
        R = s['R_mm']; r = abs(R); cx = xv + R
        semi = min(semi, r * 0.999)
        half = math.degrees(math.asin(semi / r))
        base = 180 if R > 0 else 0   # 圆心在右 → 弧在圆心左侧(180°附近)
        msp.add_arc(center=(cx, 0), radius=r, start_angle=base - half, end_angle=base + half,
                    dxfattribs={'layer': 'SURF'})
        sag = r - math.sqrt(r * r - semi * semi)
        xe = xv + (sag if R > 0 else -sag)
        ends.append((xe, semi, xe, -semi))
    # 边缘线：把每片玻璃前后两面的弧端点连起来（半口径不同时画台阶：横线 + 竖线）
    for i, s in enumerate(surfaces[:-1]):
        if s.get('stop') or not s['glass'] or ends[i] is None or ends[i + 1] is None: continue
        a, b = ends[i], ends[i + 1]
        for sign in (1, -1):
            ya, yb = a[1] * sign, b[1] * sign
            if abs(ya - yb) < 1e-6:
                msp.add_line((a[0], ya), (b[0], yb), dxfattribs={'layer': 'EDGE'})
            else:
                ymax = max(ya, yb, key=abs)
                msp.add_line((a[0], ya), (a[0], ymax), dxfattribs={'layer': 'EDGE'}) if abs(ya) < abs(ymax) else None
                msp.add_line((a[0], ymax), (b[0], ymax), dxfattribs={'layer': 'EDGE'})
                msp.add_line((b[0], ymax), (b[0], yb), dxfattribs={'layer': 'EDGE'}) if abs(yb) < abs(ymax) else None
    # 标注
    # dimdsep 不给的话 AutoCAD/渲染器会把小数点吞掉（50.00 显示成 5000）
    ov = {'dimtxt': max(1.5, total / 25), 'dimasz': max(1.2, total / 32), 'dimexo': 1.0,
          'dimdec': 2, 'dimdsep': ord('.'), 'dimzin': 0,
          'dimlfac': 1, 'dimscale': 1}   # EZDXF 样式自带 dimlfac=100（米/1:100 用），mm 图必须压回 1
    msp.add_linear_dim(base=(0, semi_max + 8), p1=(0, 0), p2=(total, 0), dimstyle='EZDXF', override=ov,
                       dxfattribs={'layer': 'DIM'}).render()
    msp.add_linear_dim(base=(-8, 0), p1=(0, -semi_max), p2=(0, semi_max), angle=90, dimstyle='EZDXF', override=ov,
                       dxfattribs={'layer': 'DIM'}).render()
    doc.saveas(path)
    _fix_image_group_order(path)


def _fix_image_group_order(path):
    """把 IMAGE 实体里的 290（裁剪模式）挪到 283 之后、360 之前 —— AutoCAD 自己就是这个顺序。
    DXF 手册把 290 列在最后，ezdxf 照手册写，但 AutoCAD 的读取器按自己的顺序来。"""
    with open(path, encoding='utf-8') as f:
        lines = f.read().split('\n')
    if len(lines) % 2:                      # 末尾换行产生的空串
        tail = lines.pop()
    else:
        tail = None
    pairs = [[lines[i], lines[i + 1]] for i in range(0, len(lines) - 1, 2)]
    out, i = [], 0
    while i < len(pairs):
        code, val = pairs[i]
        out.append(pairs[i])
        if code.strip() == '0' and val.strip() == 'IMAGE':
            j = i + 1
            rec = []
            while j < len(pairs) and pairs[j][0].strip() != '0':
                rec.append(pairs[j]); j += 1
            codes = [c.strip() for c, _ in rec]
            if '290' in codes and '283' in codes and codes.index('290') > codes.index('283'):
                pair = rec.pop(codes.index('290'))
                codes = [c.strip() for c, _ in rec]
                rec.insert(codes.index('283') + 1, pair)
            out.extend(rec)
            i = j
            continue
        i += 1
    flat = [x for pr in out for x in pr]
    if tail is not None: flat.append(tail)
    with open(path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(flat))


def build_spec(d, path, glass, title):
    nd, vd = GLASS_NDVD.get(glass.upper(), (1.51633, 64.14))
    surfs = []; k = 0; lens = 0
    for s in d['surfaces']:
        if s.get('stop'):
            surfs.append({'i': 'STO', 'R': None, 'D': s['D'], 'stop': True}); continue
        k += 1
        e = {'i': k, 'R': s['R_mm'], 'D': s['D'], 'type': '非球面' if s.get('asph') else '球面'}
        if s['glass']:
            if not s.get('cemented'): lens += 1
            e.update(nd=nd, vd=vd, glass=glass, lens=f'L{lens:02d}')
        if s['semi_mm'] is not None: e['extra'] = {'半口径': s['semi_mm']}
        if s.get('tag'): e['note'] = s['tag']
        surfs.append(e)
    spec = {'patent': title or 'diagram', 'title': title or '结构图反推', 'subtitle': f'单位 mm；玻璃全部为占位 {glass}',
            'vendors': ['HOYA', 'OHARA', 'CDGM'],
            'embodiments': [{'name': '结构图', 'extra_columns': ['半口径'], 'surfaces': surfs}]}
    json.dump(spec, open(path, 'w'), ensure_ascii=False, indent=1)


def build_seq(d, path, glass, title, fno=2.0, yan=(0.0, 10.0, 15.0)):
    """最简 CODE V 序列文件：物在无穷远，波长 d 线，口径按半口径写 CIR，玻璃占位（照用户 55EVO.seq 的写法）。"""
    gl = {'BSC7': 'BSC7_HOYA', 'N-BK7': 'NBK7_SCHOTT', 'S-BSL7': 'SBSL7_OHARA'}.get(glass.upper(), glass)
    L = ['RDM;LEN', f"TITLE '{title or 'diagram'}'", f'FNO   {fno}', 'DIM   M',
         'WL    656.3 587.6 486.1', 'REF   2', 'WTW   1 1 1',
         'YAN   ' + ' '.join(str(v) for v in yan), 'SO    0.0 0.1e19']
    for s in d['surfaces']:
        if s.get('stop'):
            L.append(f"S     0.0 {s['D']}"); L.append('  STO'); continue
        R = 0.0 if s['R_mm'] is None else s['R_mm']
        L.append(f"S     {R} {s['D']}" + (f' {gl}' if s['glass'] else ''))
        if s['semi_mm'] is not None: L.append(f"  CIR {s['semi_mm']}")
    L += ['SI    0.0 0.0', 'GO']
    open(path, 'w').write('\n'.join(L) + '\n')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('json'); ap.add_argument('-o', required=True, help='输出前缀，如 55EVO')
    ap.add_argument('--glass', default='BSC7'); ap.add_argument('--title')
    ap.add_argument('--image', help='把这张图按比例贴进 DXF（放在 DXF 同目录）')
    ap.add_argument('--seq', action='store_true', help='同时输出最简 CODE V .seq（玻璃占位，可直接 in 进 CODE V 看 layout）')
    ap.add_argument('--fno', type=float, default=2.0, help='写进 .seq 的 F 数（占位，知道就给）')
    ap.add_argument('--yan', default='0,10,15', help='写进 .seq 的视场角，逗号分隔（占位）')
    ap.add_argument('--image-keep', action='store_true', help='贴图保持原格式，不转 JPEG（AutoCAD 读 PNG 常报不可读）')
    ap.add_argument('--image-fade', type=int, default=50, help='贴图淡显百分比，0=不淡显')
    ap.add_argument('--spec', action='store_true', help='同时输出 <前缀>.spec.json 供 make_zmx.py / make_seq.py 用')
    a = ap.parse_args()
    d = json.load(open(a.json))
    x0 = d['surfaces'][0]['vertex']
    for n in clearance_clamp(d['surfaces'], d['pxmm'], x0):
        i, y0, top = n
        print(f'  面{i + 1}/{i + 2} 在 r={y0:.2f} mm 处相接（量到的机械半径 {top:.2f} mm），作图半径已截断')
    build_xlsx(d, a.o + '.xlsx', a.glass, a.title); print('写出', a.o + '.xlsx')
    build_dxf(d, a.o + '.dxf', a.image, not a.image_keep, a.image_fade); print('写出', a.o + '.dxf')
    if a.seq:
        build_seq(d, a.o + '.seq', a.glass, a.title, a.fno,
                  [float(v) for v in a.yan.split(',')]); print('写出', a.o + '.seq')
    if a.spec:
        build_spec(d, a.o + '.spec.json', a.glass, a.title); print('写出', a.o + '.spec.json')


if __name__ == '__main__':
    main()
